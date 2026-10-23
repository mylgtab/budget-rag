"""把论文 Markdown 转成可读的 Word 文档。

python scripts/md_to_docx.py --src docs/paper-draft.md --out docs/paper-draft.docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]

SKIP_HEADINGS = ("写作检查清单", "待补实验清单")
IMAGE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
STATUS = re.compile(r"\s*`?\[(已有|草稿|待写|待补|待画|待证)[^\]]*\]`?")
INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_CODE = re.compile(r"`([^`]+)`")


def set_cjk_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_runs(par, text: str) -> None:
    """处理 **粗体** 与 `代码`。"""
    text = INLINE_CODE.sub(lambda m: m.group(1), text)
    pos = 0
    for m in INLINE_BOLD.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        r = par.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def flush_table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(re.fullmatch(r":?-{2,}:?", x or "-") for x in c)]
    if not cells:
        return
    ncol = max(len(c) for c in cells)
    table = doc.add_table(rows=len(cells), cols=ncol)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(ncol):
            text = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            add_runs(par, text)
            if i == 0:
                for run in par.runs:
                    run.bold = True
            for run in par.runs:
                run.font.size = Pt(9)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default="docs/paper-draft.md")
    ap.add_argument("--out", default="docs/paper-draft.docx")
    args = ap.parse_args()

    lines = (ROOT / args.src).read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_cjk_font(doc)

    table_buf: list[str] = []
    skip = False
    title_done = False

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("## "):
            heading = STATUS.sub("", line[3:]).strip()
            if any(h in heading for h in SKIP_HEADINGS):
                skip = True
                continue
            skip = False
            doc.add_heading(heading, level=1)
            continue
        if skip:
            continue
        if line.startswith("### "):
            if table_buf:
                flush_table(doc, table_buf); table_buf = []
            doc.add_heading(STATUS.sub("", line[4:]).strip(), level=2)
            continue
        if line.startswith("# ") and not title_done:
            title_done = True
            p = doc.add_heading(STATUS.sub("", line[2:]).strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        if line.startswith("|"):
            table_buf.append(line)
            continue
        if table_buf:
            flush_table(doc, table_buf)
            table_buf = []

        if not line.strip() or line.strip() == "---":
            continue

        text = STATUS.sub("", line).strip()
        if text.startswith(">"):
            text = text.lstrip("> ").strip()
            if not text or text.startswith("状态标记"):
                continue
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.italic = True
            r.font.size = Pt(9.5)
            continue
        if text.startswith("$$") and text.endswith("$$") and len(text) > 4:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text.strip("$").strip())
            r.italic = True
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", text)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            continue
        if text.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, text[2:])
            continue
        if text.startswith("```"):
            continue
        m = IMAGE.match(text)
        if m:
            img = (ROOT / "docs" / m.group(2)).resolve()
            if img.exists():
                try:
                    doc.add_picture(str(img), width=Inches(5.6))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    doc.add_paragraph(m.group(1))
            continue
        p = doc.add_paragraph()
        add_runs(p, text)

    if table_buf:
        flush_table(doc, table_buf)

    out = ROOT / args.out
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"生成: {out}  段落={len(doc.paragraphs)}  表格={len(doc.tables)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
