"""把标注用的 xlsx 转成 utf-8 csv，交给 analyze_annotation.py。

python scripts/xlsx_to_csv.py --xlsx <文件> --out <csv>
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("需要 openpyxl：pip install openpyxl")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("空表")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        n = 0
        for r in rows[1:]:
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            w.writerow(["" if c is None else c for c in r])
            n += 1
    print(f"{args.xlsx} -> {out}（{n} 行）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
